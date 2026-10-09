# -*- coding: utf-8 -*-
"""Review the translation in Excel, then bring the edits back.

Export writes local/review/<name>.xlsx (game text inside, so it stays local):
  Тексти   key | EN | UA                            — edit UA
  Назви    key | EN | UA | Відмінки                  — edit the name, or the
           cases as "Н; Р; Д; З; О; М | Н; Р; Д; З; О; М" (singular | plural)

Import reads such a file back and merges every changed UA cell into
translation/ (cases are expanded exactly like tools/merge.py does).

Only translated lines are listed unless --all is given.

Usage:
    python tools/review.py <name> <key prefix> [<key prefix> ...] [--all]
        e.g. review.py pilot export:resource/ export:faction/ export:unit/ texts:menu.
    python tools/review.py --import local/review/<name>.xlsx
"""
import os, sys
import d4x, merge

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit('needs openpyxl:  pip install openpyxl')

EDIT = PatternFill('solid', fgColor='FFF6D5')
forms_of = merge.forms_of


def export(name, prefixes, everything=False):
    source, entries = d4x.load_source(), d4x.load_translation()
    declined = {s['base'] for s in source.values() if s.get('base')}
    plural_of = lambda k: k.replace('/texts.plural.', '/texts.') if '/texts.plural.' in k else None

    wb = openpyxl.Workbook()
    texts = wb.active
    texts.title = 'Тексти'
    names = wb.create_sheet('Назви')
    texts.append(['key', 'EN', 'UA'])
    names.append(['key', 'EN', 'UA', 'Відмінки: Н; Р; Д; З; О; М | множина'])
    n_t = n_n = 0
    for key, s in source.items():
        if not key.startswith(tuple(prefixes)) or s.get('base') or plural_of(key) in declined:
            continue
        uk = entries.get(key, {}).get('uk', '')
        if not uk and not everything:
            continue
        if key in declined:
            names.append([key, s['en'], uk, forms_of(key, entries)])
            n_n += 1
        else:
            texts.append([key, s['en'], uk])
            n_t += 1
    for ws, widths in ((texts, (34, 70, 70)), (names, (34, 30, 30, 110))):
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
            for c in row[2:]:
                c.fill = EDIT
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = 'B2'
    path = os.path.join(d4x.LOCAL, 'review', name + '.xlsx')
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print('written  %s  (%d texts, %d names)' % (os.path.relpath(path, d4x.ROOT), n_t, n_n))


def import_(path):
    source, entries = d4x.load_source(), d4x.load_translation()
    wb = openpyxl.load_workbook(path)
    changed = {}
    # columns are found by their header, so sheets written by older versions
    # of this tool (with an extra reference column) still import
    col = lambda ws, head: [str(c.value or '') for c in ws[1]].index(head)
    ws = wb['Тексти']
    ua = col(ws, 'UA')
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, uk = row[0], row[ua] or ''
        if key in source and uk != entries.get(key, {}).get('uk', ''):
            changed[key] = uk
    ws = wb['Назви']
    ua = col(ws, 'UA')
    fc = next(i for i, c in enumerate(ws[1]) if str(c.value or '').startswith('Відмінки'))
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, uk, forms = row[0], row[ua] or '', row[fc] or ''
        if key not in source:
            continue
        if forms and forms != forms_of(key, entries):
            changed.update(merge.expand(key, forms, source))
        if uk and uk != entries.get(key, {}).get('uk', ''):
            changed[key] = uk                    # an explicit name wins over the derived one
    for key, uk in changed.items():
        entries[key] = {'uk': uk, 'src': d4x.src_hash(source[key]['en'])}
    d4x.save_translation(entries, source)
    print('imported %d changed lines from %s' % (len(changed), path))


def main():
    d4x.configure_stdout()
    a = [x for x in sys.argv[1:] if x != '--all']
    if len(a) == 2 and a[0] == '--import':
        import_(a[1])
    elif len(a) >= 2:
        export(a[0], a[1:], '--all' in sys.argv)
    else:
        sys.exit(__doc__)


if __name__ == '__main__':
    main()
